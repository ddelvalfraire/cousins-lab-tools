"""Reads the researcher's mass-spec workbook (`Raw data-mass spec.xlsx`).

Sheet `RawDataAcquisition67Data` holds `Time, Mass 32, Mass 44, Mass 45` every 0.5 s.
Time is seconds from the start of that acquisition; the sheet records no clock time.
Which program produced the 0.5 s series is unknown.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .contract import Clock, Column

RAW_SHEET = "RawDataAcquisition67Data"
RAW_HEADER = ("Time", "Mass 32", "Mass 44", "Mass 45")

MS_IMPORT_COLUMNS = (
    Column("seq", "int64", description="Row number within the sheet's data"),
    Column("t_elapsed_s", "float64", "s", "Workbook Time column: seconds since that acquisition started"),
    Column("m32", "float64", "legacy units (unverified)", "Mass 32; constant -404.404 in this file"),
    Column("m44", "float64", "mV per workbook calibration", "Mass 44 (CO2)"),
    Column("m45", "float64", "legacy units (unverified)", "Mass 45"),
)
MS_IMPORT_CLOCKS = (Clock("t_elapsed_s", "elapsed", 1.0, "Seconds since the MS acquisition started; no clock time recorded"),)


@dataclass(frozen=True)
class MassSpecSeries:
    t_elapsed_s: list[float]
    m32: list[float]
    m44: list[float]
    m45: list[float]


def read_raw_sheet(path: Path, sheet: str = RAW_SHEET) -> MassSpecSeries:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb[sheet].iter_rows(min_col=1, max_col=4, values_only=True)
        header = next(rows)
        if tuple(header) != RAW_HEADER:
            raise ValueError(f"{sheet} header is {header}, expected {RAW_HEADER}")
        t, m32, m44, m45 = [], [], [], []
        for row in rows:
            if row[0] is None:
                break
            t.append(float(row[0]))
            m32.append(float(row[1]))
            m44.append(float(row[2]))
            m45.append(float(row[3]))
        return MassSpecSeries(t, m32, m44, m45)
    finally:
        wb.close()


def read_sample_sheet(path: Path, sheet: str) -> list[tuple[float, float]]:
    """(T(s), ubar CO2) pairs from a researcher's per-sample sheet such as `150-pl1`."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        out = []
        for row in wb[sheet].iter_rows(min_row=2, min_col=1, max_col=5, values_only=True):
            if row[0] is None:
                break
            out.append((float(row[0]), float(row[4])))
        return out
    finally:
        wb.close()
